# app/fiscal/apuracao_logic.py
import logging
import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from typing import Optional, Any, List, Union
import FreeSimpleGUI as sg

def _encontrar_celula(ws: Worksheet, texto_procurado: str) -> Optional[tuple[int, int]]:
    """
    Encontra as coordenadas (linha, coluna) de uma célula baseada no texto exato.
    """
    if texto_procurado is None: return None
    texto_limpo = str(texto_procurado).strip()

    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and str(cell.value).strip() == texto_limpo:
                return (cell.row, cell.column)
    return None

def _escrever_valor_adjacente(ws: Worksheet, label_texto: str, valor: Any, offset_col: int = 1):
    """
    Escreve o valor na célula ao lado (offset_col=1) da célula onde encontrou o 'label_texto'.
    """
    coordenadas = _encontrar_celula(ws, label_texto)
    if coordenadas:
        (row, col) = coordenadas
        try:
            celula_alvo = ws.cell(row=row, column=col + offset_col)
            celula_alvo.value = valor
        except Exception as e:
            logging.warning(f"Erro ao escrever valor para '{label_texto}': {e}")
    else:
        # Loga aviso mas não para o processo
        logging.warning(f"RÓTULO NÃO ENCONTRADO NO EXCEL: '{label_texto}'")

def somar_por_cfop(df: pd.DataFrame, cfops: Union[str, List[str]], coluna_valor: str) -> float:
    """
    Filtra o DataFrame por uma lista de CFOPs e soma a coluna especificada.
    """
    if df.empty: return 0.0

    lista_cfops = [cfops] if isinstance(cfops, str) else cfops
    # Converte para string para garantir match
    mask = df['CFOP (SPED)'].astype(str).isin(lista_cfops)

    val = df[mask][coluna_valor].sum()
    return float(val)

def preencher_template_apuracao(template_path: Path, df_entradas: pd.DataFrame, df_saidas: pd.DataFrame) -> None:
    """
    Preenche o template de apuração com base em uma lista de regras estruturadas.
    """
    logging.info(f"Iniciando preenchimento da apuração: {template_path.name}")

    try:
        wb = load_workbook(template_path)
        ws = wb["Apuracao"] if "Apuracao" in wb.sheetnames else wb.active
    except Exception as e:
        logging.error(f"Erro ao carregar arquivo Excel: {e}"); raise

    # ==============================================================================
    # ESTRUTURA DE REGRAS PARA APURAÇÃO
    # Define como cada campo do Excel deve ser calculado.
    # 'label': O texto exato a ser encontrado na planilha.
    # 'tipo_df': 'entradas' ou 'saidas', para saber qual DataFrame usar.
    # 'coluna': A coluna do DataFrame a ser somada.
    # 'tipo_calculo': 'total' (soma a coluna inteira) ou 'cfop' (filtra por CFOPs antes de somar).
    # 'cfops': Lista de CFOPs a serem usados quando tipo_calculo é 'cfop'.
    # ==============================================================================
    regras_apuracao = [
        # --- 1. ENTRADAS (TOTAIS) ---
        {'label': "Total Contabil Entradas", 'tipo_df': 'entradas', 'coluna': 'Total Operação', 'tipo_calculo': 'total'},
        {'label': "Total Base ICMS Entradas", 'tipo_df': 'entradas', 'coluna': 'Base de Cálculo ICMS', 'tipo_calculo': 'total'},
        {'label': "Total Credito ICMS", 'tipo_df': 'entradas', 'coluna': 'Total ICMS', 'tipo_calculo': 'total'},
        {'label': "Total Credito IPI", 'tipo_df': 'entradas', 'coluna': 'Total IPI', 'tipo_calculo': 'total'},
        {'label': "Total Base ST Entradas", 'tipo_df': 'entradas', 'coluna': 'Base de Cálculo ICMS ST', 'tipo_calculo': 'total'},
        {'label': "Total ICMS ST Entradas", 'tipo_df': 'entradas', 'coluna': 'Total ICMS ST', 'tipo_calculo': 'total'},

        # --- 2. SAÍDAS (TOTAIS) ---
        {'label': "Total Contabil Saidas", 'tipo_df': 'saidas', 'coluna': 'Total Operação', 'tipo_calculo': 'total'},
        {'label': "Total Base ICMS Saidas", 'tipo_df': 'saidas', 'coluna': 'Base de Cálculo ICMS', 'tipo_calculo': 'total'},
        {'label': "Total Debito ICMS", 'tipo_df': 'saidas', 'coluna': 'Total ICMS', 'tipo_calculo': 'total'},
        {'label': "Total Debito IPI", 'tipo_df': 'saidas', 'coluna': 'Total IPI', 'tipo_calculo': 'total'},
        {'label': "Total Base ST Saidas", 'tipo_df': 'saidas', 'coluna': 'Base de Cálculo ICMS ST', 'tipo_calculo': 'total'},
        {'label': "Total ICMS ST Saidas", 'tipo_df': 'saidas', 'coluna': 'Total ICMS ST', 'tipo_calculo': 'total'},

        # --- 3. DETALHAMENTO SAÍDAS (POR CFOP) ---
        {'label': "Venda Producao (5101/6101)", 'tipo_df': 'saidas', 'coluna': 'Total Operação', 'tipo_calculo': 'cfop', 'cfops': ['5101', '6101', '5103', '6103']},
        {'label': "Revenda Mercadoria (5102/6102)", 'tipo_df': 'saidas', 'coluna': 'Total Operação', 'tipo_calculo': 'cfop', 'cfops': ['5102', '6102']},
        {'label': "Venda c/ ST (5403/5405)", 'tipo_df': 'saidas', 'coluna': 'Total Operação', 'tipo_calculo': 'cfop', 'cfops': ['5403', '5405', '6403', '6404']},
        {'label': "Outras Saidas (Remessas/Bonif)", 'tipo_df': 'saidas', 'coluna': 'Total Operação', 'tipo_calculo': 'cfop', 'cfops': ['5910', '6910', '5949', '6949']},

        # --- 4. DETALHAMENTO ENTRADAS (POR CFOP) ---
        {'label': "Compra p/ Industrializacao (1101/2101)", 'tipo_df': 'entradas', 'coluna': 'Total Operação', 'tipo_calculo': 'cfop', 'cfops': ['1101', '2101']},
        {'label': "Compra p/ Comercializacao (1102/2102)", 'tipo_df': 'entradas', 'coluna': 'Total Operação', 'tipo_calculo': 'cfop', 'cfops': ['1102', '2102']},
        {'label': "Uso e Consumo (1556/2556)", 'tipo_df': 'entradas', 'coluna': 'Total Operação', 'tipo_calculo': 'cfop', 'cfops': ['1556', '2556']},
        {'label': "Ativo Imobilizado (1551/2551)", 'tipo_df': 'entradas', 'coluna': 'Total Operação', 'tipo_calculo': 'cfop', 'cfops': ['1551', '2551']},
        {'label': "Compra c/ ST (1403/2403)", 'tipo_df': 'entradas', 'coluna': 'Total Operação', 'tipo_calculo': 'cfop', 'cfops': ['1403', '2403']},
        {'label': "Devolucao de Venda (1202/2202)", 'tipo_df': 'entradas', 'coluna': 'Total Operação', 'tipo_calculo': 'cfop', 'cfops': ['1202', '2202']},

        # --- 5. IMPOSTOS ESPECÍFICOS (POR CFOP) ---
        {'label': "ICMS da Producao Propria", 'tipo_df': 'saidas', 'coluna': 'Total ICMS', 'tipo_calculo': 'cfop', 'cfops': ['5101', '6101']},
        {'label': "ICMS da Revenda", 'tipo_df': 'saidas', 'coluna': 'Total ICMS', 'tipo_calculo': 'cfop', 'cfops': ['5102', '6102']},
        {'label': "ICMS ST da Revenda (5405)", 'tipo_df': 'saidas', 'coluna': 'Total ICMS ST', 'tipo_calculo': 'cfop', 'cfops': ['5405']},
    ]

    # Processamento genérico das regras
    for regra in regras_apuracao:
        df_alvo = df_entradas if regra['tipo_df'] == 'entradas' else df_saidas
        valor = 0.0

        if not df_alvo.empty:
            if regra['tipo_calculo'] == 'total':
                valor = df_alvo[regra['coluna']].sum()
            elif regra['tipo_calculo'] == 'cfop':
                valor = somar_por_cfop(df_alvo, regra.get('cfops', []), regra['coluna'])

        _escrever_valor_adjacente(ws, regra['label'], valor)

    try:
        wb.save(template_path)
        logging.info(f"Arquivo '{template_path.name}' atualizado com sucesso.")
    except Exception as e:
        sg.popup_error(f"Erro ao salvar apuração.\nVerifique se o arquivo '{template_path.name}' não está aberto.\nErro: {e}")
        raise
